"""Extração de data e assunto de um arquivo."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

_PALAVRAS_IGNORADAS = {"de", "da", "do", "e", "a", "o", "para", "com", "final", "copia", "cópia"}
_SEPARADORES = re.compile(r"[_\-\s]+")
_APENAS_NUMEROS = re.compile(r"^\d+$")


def obter_data(caminho: Path) -> datetime:
    """Data de referência do arquivo (data de modificação)."""
    return datetime.fromtimestamp(caminho.stat().st_mtime)


def _assunto_por_metadados(caminho: Path) -> str | None:
    """Tenta extrair um título a partir dos metadados do documento (best-effort)."""
    extensao = caminho.suffix.lower()
    try:
        if extensao == ".pdf":
            from pypdf import PdfReader

            titulo = PdfReader(str(caminho)).metadata.title
            return titulo.strip() if titulo else None
        if extensao == ".docx":
            from docx import Document

            titulo = Document(str(caminho)).core_properties.title
            return titulo.strip() if titulo else None
        if extensao == ".xlsx":
            from openpyxl import load_workbook

            titulo = load_workbook(str(caminho), read_only=True).properties.title
            return titulo.strip() if titulo else None
    except Exception:
        return None
    return None


def _assunto_por_nome(caminho: Path) -> str:
    """Deriva um assunto legível a partir do nome do arquivo, ignorando datas/números."""
    partes = _SEPARADORES.split(caminho.stem)
    palavras = [
        parte
        for parte in partes
        if parte and not _APENAS_NUMEROS.match(parte) and parte.lower() not in _PALAVRAS_IGNORADAS
    ]
    if not palavras:
        return "Geral"
    return " ".join(palavras).title()


def obter_assunto(caminho: Path) -> str:
    """Obtém o assunto do arquivo: primeiro pelos metadados, depois pelo nome."""
    assunto = _assunto_por_metadados(caminho)
    if assunto:
        return assunto
    return _assunto_por_nome(caminho)
